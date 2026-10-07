#!/usr/bin/env python3
"""Legge source.xlsx e genera app/src/main/assets/seed.js (dati iniziali dell'app).
Uso:  python3 tools/build_seed.py   (dalla radice del progetto)"""
import openpyxl, json, re, sys, os
from openpyxl.utils import get_column_letter as L

SRC = os.path.join(os.path.dirname(__file__), '..', 'source.xlsx')
OUT = os.path.join(os.path.dirname(__file__), '..', 'app', 'src', 'main', 'assets', 'seed.js')
wv = openpyxl.load_workbook(SRC, data_only=True)

def ym(i):  # indice mese 0 = aprile 2023
    n = 3 + i  # 0-based month da gennaio 2023
    return '%04d-%02d' % (2023 + n // 12, n % 12 + 1)

def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)

def clean(s):
    return re.sub(r'\s+', ' ', str(s)).strip()

movs, seq = [], [0]
def add(t, m, w, s, c, a, k=None):
    seq[0] += 1
    movs.append({'id': 's%d' % seq[0], 't': t, 'm': m, 'd': 0, 'w': w, 's': s,
                 'c': c, 'a': round(float(a), 2), 'k': k, 'n': '', 'u': 0})

# ---------- conti ----------
ACC = {}  # nome -> dict
def acc(name, owner, kind):
    if name not in ACC:
        ACC[name] = {'id': 'a%d' % (len(ACC) + 1), 'name': name, 'owner': owner, 'kind': kind, 'base': 0, 'asof': None}
    return ACC[name]['id']
for n, o, k in [('Conto Poste', 'C', 'bank'), ('Contanti Pao', 'P', 'cash'), ('Contanti Sas Prendas', 'P', 'cash'),
                ('Conto Unicredit Pao', 'P', 'bank'), ('Contanti Simo', 'S', 'cash'), ('Conto BDS Simo', 'S', 'bank'),
                ('Conto ING Simo', 'S', 'bank'), ('Paypal Simo', 'S', 'bank'), ('Regalo', 'C', 'cash'),
                ('Conto BDS Pao Simo', 'C', 'bank'), ('Libretto poste pao', 'P', 'bank')]:
    acc(n, o, k)

# ---------- spese casa ----------
def house(sheet, first_month_idx, aprile_solo_importi=False):
    w = wv[sheet]
    tot_row = w.max_row + 1  # i mesi lunghi superano la riga "Totale": si legge tutto, le celle senza tipologia sono totali
    hdr = [c.column for c in w[2] if c.value == 'Tipologia']
    n0 = len(movs)
    idx = first_month_idx
    if aprile_solo_importi:  # colonna B: aprile 2023, solo importi
        tr = next(r for r in range(1, 100) if str(w.cell(r, 1).value or '').strip() == 'Totale')
        for r in range(3, tr):
            v = w.cell(r, 2).value
            if num(v) and v: add('o', ym(idx), 'C', 'c', 'Spesa casa (senza dettaglio)', v)
        idx += 1
    for col in hdr:
        m = ym(idx); idx += 1
        for r in range(3, tot_row):
            t, a, k = w.cell(r, col).value, w.cell(r, col + 1).value, w.cell(r, col + 2).value
            if t is None or not num(a) or str(t).startswith('-'): continue
            k = clean(k) if k else None
            owner = ACC[k]['owner'] if k in ACC else 'C'
            add('o', m, owner, 'c', clean(t), a, ACC[k]['id'] if k in ACC else None)
    return idx, tot_row, n0
i1, _, n1 = house('Expenses Archivio', 0, True)
i2, _, n2 = house('Expenses da aprile 2025', 24)
assert i1 == 24 and i2 == 42, (i1, i2)

# controllo: somma per mese vs cella totale del foglio (importo senza tipologia)
w = wv['Expenses da aprile 2025']
hdr = [c.column for c in w[2] if c.value == 'Tipologia']
bad = 0
for j, col in enumerate(hdr):
    m = ym(24 + j)
    mine = round(sum(x['a'] for x in movs if x['m'] == m and x['s'] == 'c'), 2)
    tots = [w.cell(r, col + 1).value for r in range(3, w.max_row + 1)
            if w.cell(r, col).value is None and num(w.cell(r, col + 1).value) and w.cell(r, col + 1).value > 100]
    if tots and abs(tots[-1] - mine) > 0.02:
        bad += 1; print('DIFF', m, mine, tots[-1])
print('controllo totali spese casa: %d differenze' % bad)

# ---------- spese personali ----------
wf = openpyxl.load_workbook(SRC, data_only=False)
def personal(sheet, who, ncols):
    w, f = wv[sheet], wf[sheet]
    tot = next(r for r in range(1, 100) if str(w.cell(r, 1).value or '').strip() == 'Totale')
    for j in range(ncols):
        m = ym(j); col = 2 + j
        cells = []
        for r in range(3, w.max_row + 1):
            v, fv = w.cell(r, col).value, f.cell(r, col).value
            if not num(v) or not v: continue
            if isinstance(fv, str) and fv.startswith('=') : continue   # formula di totale
            cells.append((r, v))
        # riga "Totale" con valore fisso coerente con la somma sopra: è un totale, non una spesa
        above = [v for r, v in cells if r < tot]
        below = [v for r, v in cells if r > tot]
        tv = [v for r, v in cells if r == tot]
        if tv and not below and abs(tv[0] - sum(above)) < 0.02:
            cells = [(r, v) for r, v in cells if r != tot]
        for r, v in cells: add('o', m, who, 'p', 'Spese personali', v)
personal('Expenses Simo', 'S', 19)
personal('Expenses Pao', 'P', 43)

# ---------- entrate ----------
w = wv['Income']
rows = {2: 'C', 3: 'S', 4: 'S', 5: 'P', 6: 'P', 7: 'P', 8: 'C', 9: 'P'}
for r, who in rows.items():
    cat = clean(w.cell(r, 1).value)
    for j in range(43):
        v = w.cell(r, 2 + j).value
        if num(v) and v: add('i', ym(j), who, 'c', cat, v)

# ---------- resoconti (saldi conti + storico patrimonio) ----------
def resoconto(sheet, who, alias, ncols, total_row):
    w = wv[sheet]
    tot = [w.cell(total_row, 2 + j).value for j in range(ncols)]
    last = max(j for j, v in enumerate(tot) if num(v) and v != 0)
    hist = {ym(j): round(v, 2) for j, v in enumerate(tot) if num(v) and v != 0}
    r = 3
    while r < total_row:
        name = clean(w.cell(r, 1).value or '')
        if name:
            name = alias.get(name, name)
            if name not in ACC:
                kind = 'cash' if 'ontanti' in name else 'bank'
                acc(name, who if who != 'X' else 'C', kind)
            v = w.cell(r, 2 + last).value
            if num(v):
                ACC[name]['base'] = round(float(v), 2); ACC[name]['asof'] = ym(last)
        r += 1
    return hist
hist = {}
hist['S'] = resoconto('Resoconto Simo', 'S', {'Conto Paypal': 'Paypal Simo', 'Conto BDS': 'Conto BDS Simo'}, 19, 10)
hist['P'] = resoconto('Resoconto Pao', 'P', {'Libretto postale': 'Libretto poste pao', 'Conto bds': 'Conto bds', 'Conto ing': 'Conto ing'}, 43, 12)
for n in ('Conto bds', 'Conto ing'):
    ACC[n]['owner'] = 'C'
w = wv['Resoconto Simo']
hist['C'] = {ym(j): round(w.cell(11, 2 + j).value, 2) for j in range(19) if num(w.cell(11, 2 + j).value)}

seed = {'v': 1, 'users': [{'id': 'S', 'name': 'Simo'}, {'id': 'P', 'name': 'Paola'}],
        'accounts': list(ACC.values()), 'movs': movs, 'hist': hist, 'lastMonth': ym(42)}
with open(OUT, 'w', encoding='utf-8') as f:
    f.write('window.SEED=' + json.dumps(seed, ensure_ascii=False, separators=(',', ':')) + ';')
print('movimenti:', len(movs), '| conti:', len(ACC), '| bytes:', os.path.getsize(OUT))
for a in ACC.values(): print(' ', a['id'], a['name'], a['owner'], a['kind'], a['base'], a['asof'])
